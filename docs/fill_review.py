# -*- coding: utf-8 -*-
"""把文献综述内容填入学院模板。

模板格式规范（模板自身注明）：
  一级标题：黑体 小三(15pt) 加粗 1.5倍行距
  二级标题：宋体 四号(14pt) 加粗 1.5倍行距
  正文    ：中文宋体 / 英文 Times New Roman，小四(12pt)，27磅行距，首行缩进2字符
  参考文献：另起一页；标题黑体小三加粗居中；条目 五号(10.5pt)，27磅行距

引用编号：正文写 ``[[key]]``，本脚本按 GB/T 7714 顺序编码制（首次出现顺序）
自动替换为 ``[n]`` 并生成对应顺序的参考文献表，杜绝人工编号错位。
``{{REFSTATS}}`` 占位符由脚本按实际引用文献的类型分布自动填写。
"""
import re
import sys
import os
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from review_content import TITLE, BODY
from refs import REFDB

TEMPLATE = r"C:\Users\tyc\Desktop\毕业设计（论文）材料模板\《毕业设计（论文）》文献综述.docx"
OUT = r"C:\Users\tyc\Desktop\《毕业设计（论文）》文献综述-填写版.docx"

HEI, SONG, TNR = "黑体", "宋体", "Times New Roman"

KEY_RE = re.compile(r"\[\[([A-Za-z0-9_]+)\]\]")

# ─────────────────────────────────────────────────────────────
# 1. 计算引用编号（顺序编码制）
# ─────────────────────────────────────────────────────────────
BODY_TEXT = "\n".join(t for _, t in BODY if isinstance(t, str))

NUM = {}
ORDER = []
for m in KEY_RE.finditer(BODY_TEXT):
    k = m.group(1)
    if k not in REFDB:
        raise KeyError("未登记的文献键: %s" % k)
    if k not in NUM:
        NUM[k] = len(NUM) + 1
        ORDER.append(k)

REFS = [REFDB[k] for k in ORDER]

# ─────────────────────────────────────────────────────────────
# 2. 自动生成文献构成统计（替代易失真的手写数字）
# ─────────────────────────────────────────────────────────────
TYPE_NAME = {"M": "专著", "J": "期刊论文", "C": "会议论文",
             "D": "学位论文", "R": "技术报告", "EB/OL": "电子资源"}
TYPE_ORDER = ["M", "J", "C", "D", "R", "EB/OL"]

_tc = Counter()
for r in REFS:
    m = re.search(r"\[(M|J|C|D|R|EB/OL)\]", r)
    if not m:
        raise ValueError("文献缺少 GB/T 7714 类型标识: %s" % r[:60])
    _tc[m.group(1)] += 1

_cn = sum(1 for r in REFS if re.search(r"[\u4e00-\u9fff]", r[:14]))
_yrs = []
for r in REFS:
    ys = re.findall(r"(19\d\d|20\d\d)", r)
    if ys:
        _yrs.append(max(int(y) for y in ys))

_parts = "、".join("%s %d %s" % (TYPE_NAME[t], _tc[t], "部" if t == "M" else
                                ("份" if t == "R" else ("项" if t == "EB/OL" else "篇")))
                   for t in TYPE_ORDER if _tc.get(t))
REFSTATS = ("本文所引文献共 %d 篇，其中%s；中文文献 %d 篇，时间跨度自 %d 年至 %d 年。"
            % (len(REFS), _parts, _cn, min(_yrs), max(_yrs)))


def sub_keys(s):
    return KEY_RE.sub(lambda m: "[%d]" % NUM[m.group(1)], s).replace("{{REFSTATS}}", REFSTATS)


print("引用文献 %d 条（编号 1..%d）" % (len(REFS), len(REFS)))
print(REFSTATS)

# ─────────────────────────────────────────────────────────────
# 3. Word 渲染
# ─────────────────────────────────────────────────────────────
def set_font(run, cn=SONG, en=TNR, size=12, bold=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)


CN_PUNCT = "，。、；：？！（）《》【】“”—…·"


def split_cn_en(text):
    """按中/非中切分：中文走宋体(或黑体)，西文与数字走 Times New Roman。"""
    segs, cur, cur_is_cn = [], "", None
    for ch in text:
        is_cn = "\u4e00" <= ch <= "\u9fff" or ch in CN_PUNCT
        if cur_is_cn is None:
            cur_is_cn = is_cn
        if is_cn != cur_is_cn:
            segs.append((cur, cur_is_cn))
            cur, cur_is_cn = "", is_cn
        cur += ch
    if cur:
        segs.append((cur, cur_is_cn))
    return segs


doc = docx.Document(TEMPLATE)

# 封面：填入题目
for p in doc.paragraphs[:17]:
    if p.text.strip().startswith("题目"):
        p.runs[0].text = "题目：  " + TITLE
        set_font(p.runs[0], size=16)

# 删除模板正文示例段（“1一级标题…”起至示例参考文献）
paras = doc.paragraphs
start = None
for i, p in enumerate(paras):
    if p.text.strip().startswith("1一级标题"):
        start = i
        break
assert start is not None, "未找到模板正文起点"
for p in paras[start:]:
    p._element.getparent().remove(p._element)

sectPr = doc.element.body.find(qn("w:sectPr"))   # 节属性必须留在文档末尾
anchor = sectPr


def para(text, *, cn=SONG, size=12, bold=False, align=None, line=27,
         line_rule="exact", first_indent=None, before=None, after=None,
         page_break=False):
    """新建段落并插到 sectPr 之前；返回该段落对象。"""
    p = doc.add_paragraph()
    if anchor is not None:
        anchor.addprevious(p._element)
    pf = p.paragraph_format
    if line_rule == "exact":
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if align is not None:
        pf.alignment = align
    if first_indent is not None:
        pf.first_line_indent = first_indent
        # 同时写入 firstLineChars="200"（Word 的“2 字符”单位），比固定 twips 稳
        ind = p._p.get_or_add_pPr().find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            p._p.get_or_add_pPr().append(ind)
        ind.set(qn("w:firstLineChars"), "200")
    if before is not None:
        pf.space_before = before
    if after is not None:
        pf.space_after = after
    if page_break:
        pf.page_break_before = True
    for seg, is_cn in split_cn_en(text):
        r = p.add_run(seg)
        # eastAsia 统一用 cn（标题黑体 / 正文宋体），西文与数字一律 TNR
        set_font(r, cn=cn, en=TNR, size=size, bold=bold)
    return p


TWO_CHAR = Pt(24)   # 小四(12pt) × 2 字符

# 正文
for kind, raw in BODY:
    text = sub_keys(raw)
    if kind == "h1":
        para(text, cn=HEI, size=15, bold=True, line_rule="1.5",
             before=Pt(12), after=Pt(6))
    elif kind == "h2":
        para(text, size=14, bold=True, line_rule="1.5",
             before=Pt(8), after=Pt(4))
    else:
        para(text, size=12, line=27, first_indent=TWO_CHAR)

# 参考文献（另起一页）
para("参考文献", cn=HEI, size=15, bold=True, line_rule="1.5",
     align=WD_ALIGN_PARAGRAPH.CENTER, page_break=True)
for i, ref in enumerate(REFS, 1):
    para("[%d] %s" % (i, ref), size=10.5, line=27, after=Pt(2))

doc.save(OUT)
print("saved:", OUT)
print("正文块:", len(BODY), "| 参考文献:", len(REFS))
